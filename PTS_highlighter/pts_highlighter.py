"""Prepare piping execution views while preserving the source Excel workbook."""

import argparse
from collections import Counter
from copy import copy
from datetime import date, datetime
from difflib import SequenceMatcher
import hashlib
import json
from pathlib import Path
import re
import unicodedata

import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.filters import AutoFilter


DISCIPLINES = ("Piping", "Hangers", "Valves", "ALL")
STATUS_TITLE = "PTS MILESTONE STATUS"
CALCULATOR_NAME = "_PTS_Weeks_Calculator"
COLOURS = {
    "Lot General Piping Pipe Material": "FFE2F0D9",
    "Lot General Piping Engineering": "FFD9EAF7",
    "Lot General Piping Hangers": "FFFCE4D6",
    "Lot General Piping Valves": "FFFFF2CC",
    "Lot 5 - Erection General Piping": "FFE4DFEC",
}
PROCUREMENT = (
    "TSD / BANF RFQ", "Tender Due Date",
    "Technical & Commercial Comparison Pre-TCC", "TSD PO",
    "Technical & Commercial Comparison TCC", "Purchase Order",
    "Ready for Shipment", "Shipment", "DTS (T)", "DTS (S)",
)
ENGINEERING = (
    "BOM / BOQ-2 - Issue for for PO General Piping Material",
    "BOM and ISO and arrangement Drawings for PO General Piping Erection",
)
ERECTION = (
    "BOQ 1 (Doc's for RfQ) - Issued from Engineering",
    "BOQ 2 (Doc's for PO) - Issued from Engineering",
    "Contractor On Site Mobilisation Completed",
)
RULES = {
    "Lot General Piping Pipe Material": (PROCUREMENT, ("Piping",)),
    "Lot General Piping Engineering": (ENGINEERING, ("Piping", "Hangers")),
    "Lot General Piping Hangers": (
        ("TSD / BANF RFQ", "Purchase Order", "Ready for Shipment", "Shipment",
         "DTS (T)", "DTS (S)"), ("Hangers",)
    ),
    "Lot General Piping Valves": (PROCUREMENT, ("Valves",)),
    "Lot 5 - Erection General Piping": (ERECTION, ("Piping",)),
}
HEADER_ALIASES = {
    "Activity ID": ("activity id", "activity identifier", "activity code", "actvity id"),
    "Activity Name": ("activity name", "activity description", "actvity name"),
    "Start": ("start", "start date", "current start", "current start date"),
    "Finish": ("finish", "finish date", "current finish", "current finish date"),
}
CLASS_ALIASES = {
    "piping": "Piping", "pipe": "Piping", "pipes": "Piping",
    "hanger": "Hangers", "hangers": "Hangers", "valve": "Valves",
    "valves": "Valves", "all": "ALL",
}


def normalize(value: object) -> str:
    text = unicodedata.normalize("NFKD", str(value or "")).casefold()
    text = text.replace("&", " and ").replace("'", "").replace("\u2019", "")
    words = re.findall(r"[a-z0-9]+", text)
    return " ".join(word for index, word in enumerate(words)
                    if word != "the" and not (index and word == words[index - 1]))


def close_words(actual: str, expected: str) -> bool:
    actual_words, expected_words = actual.split(), expected.split()
    if len(actual_words) != len(expected_words):
        return False
    differences = [(left, right) for left, right in zip(actual_words, expected_words)
                   if left != right]
    return len(differences) <= 1 and all(
        min(len(left), len(right)) >= 4
        and SequenceMatcher(None, left, right).ratio() >= 0.85
        for left, right in differences
    )


def classify(value: object) -> dict:
    text = normalize(value)
    for family, (milestones, disciplines) in RULES.items():
        prefix = normalize(family)
        count = len(prefix.split())
        tokens = text.split()
        if not close_words(" ".join(tokens[:count]), prefix):
            continue
        suffix = " ".join(tokens[count:])
        matches = [milestone for milestone in milestones
                   if close_words(suffix, normalize(milestone))]
        if len(matches) == 1:
            return {"family": family, "milestone": matches[0],
                    "disciplines": disciplines, "ambiguous": False}
        ambiguous = not suffix or any(
            SequenceMatcher(None, suffix, normalize(milestone)).ratio() >= 0.68
            for milestone in milestones
        )
        return {"family": family, "milestone": None,
                "disciplines": (), "ambiguous": ambiguous}
    return {"family": None, "milestone": None,
            "disciplines": (), "ambiguous": False}


def clean_date(value: object) -> tuple[object, bool]:
    if value is None or value == "":
        return value, True
    if isinstance(value, (datetime, date)):
        return datetime(value.year, value.month, value.day), True
    if not isinstance(value, str) or value.startswith("="):
        return value, False
    matches = re.findall(
        r"(?<!\d)(\d{1,2})\s*[./-]\s*(\d{1,2})\s*[./-]\s*(\d{4})(?!\d)", value
    )
    dates = set()
    for day, month, year in matches:
        try:
            dates.add(datetime(int(year), int(month), int(day)))
        except ValueError:
            pass
    return (dates.pop(), True) if len(dates) == 1 else (value, False)


def select_sheet(book, requested: str | None = None):
    visible = [sheet for sheet in book if sheet.sheet_state == "visible"
               and sheet["A1"].value != STATUS_TITLE]
    if requested:
        choices = [sheet for sheet in visible if sheet.title == requested]
        if not choices:
            raise ValueError(f"Visible worksheet not found: {requested}")
        return choices[0]
    if len(visible) == 1:
        return visible[0]
    dated = []
    for sheet in visible:
        match = re.search(r"\d{1,2}[A-Za-z]{3}\d{4}", sheet.title)
        if not match:
            raise ValueError("Cannot identify latest visible worksheet; use --sheet")
        dated.append((datetime.strptime(match.group(), "%d%b%Y"), sheet))
    dated.sort(key=lambda item: item[0], reverse=True)
    if not dated or (len(dated) > 1 and dated[0][0] == dated[1][0]):
        raise ValueError("Latest visible worksheet is ambiguous; use --sheet")
    return dated[0][1]


def detect_headers(sheet) -> tuple[int, dict, dict]:
    candidates = []
    for row_number in range(1, min(sheet.max_row, 100) + 1):
        values = {column: normalize(cell.value)
                  for (row, column), cell in sheet._cells.items()
                  if row == row_number and cell.value is not None}
        found = {}
        for name, aliases in HEADER_ALIASES.items():
            matches = [(aliases.index(value), column) for column, value in values.items()
                       if value in aliases]
            if matches:
                matches.sort()
                if len(matches) > 1 and matches[0][0] == matches[1][0]:
                    raise ValueError(f"Ambiguous {name} header in row {row_number}")
                found[name] = matches[0][1]
        if found:
            candidates.append((len(found), row_number, found, values))
    if not candidates:
        raise ValueError("Required headers not found: Activity ID, Activity Name, Start, Finish")
    candidates.sort(key=lambda item: (-item[0], item[1]))
    score, row_number, found, values = candidates[0]
    if score != 4:
        raise ValueError("Required headers not found: " + ", ".join(
            name for name in HEADER_ALIASES if name not in found))
    classifications = {}
    for column, value in values.items():
        if value in CLASS_ALIASES:
            name = CLASS_ALIASES[value]
            if name in classifications:
                raise ValueError(f"Duplicate classification column: {name}")
            classifications[name] = column
    return row_number, found, classifications


def sheet_fingerprint(sheet) -> str:
    digest = hashlib.sha256()
    digest.update(repr((sheet.title, sheet.sheet_state, sheet.freeze_panes,
                        str(sheet.auto_filter.ref))).encode())
    for key, cell in sorted(sheet._cells.items()):
        digest.update(repr((key, cell.value, cell.data_type, cell._style)).encode())
    return digest.hexdigest()


def relocate_columns(sheet, finish: int, classifications: dict) -> dict:
    """Move sparse cells without expanding Excel's potentially inflated used range."""
    removed = set(classifications.values())
    if any(column <= finish for column in removed):
        raise ValueError("Classification columns before Finish are not supported yet")
    mapping = {
        column: column + (4 if column > finish else 0)
        - sum(old < column for old in removed)
        for column in range(1, sheet.max_column + 1) if column not in removed
    }
    if max(mapping.values()) > 16384:
        raise ValueError("Moving columns would exceed Excel's column limit")
    if sheet.merged_cells.ranges or sheet.tables or sheet.conditional_formatting:
        raise ValueError("Merged cells, tables or conditional formatting need a project-specific review")
    if any(cell.data_type == "f" and column not in removed
           for (_, column), cell in sheet._cells.items()):
        raise ValueError("Original schedule formulas require reference-aware column relocation")
    moved = {}
    for (row, column), cell in list(sheet._cells.items()):
        if column not in removed:
            cell.column = mapping[column]
            moved[(row, cell.column)] = cell
    sheet._cells = moved
    dimensions = list(sheet.column_dimensions.values())
    sheet.column_dimensions.clear()
    for dimension in dimensions:
        first = dimension.min or openpyxl.utils.column_index_from_string(dimension.index)
        for column in range(first, (dimension.max or first) + 1):
            if column in mapping:
                target = mapping[column]
                result = copy(dimension)
                result.index = get_column_letter(target)
                result.min = result.max = target
                sheet.column_dimensions[result.index] = result
    return mapping


def schedule_sort_key(finish: object, name: object) -> tuple:
    return (finish.toordinal() if isinstance(finish, (date, datetime)) else float("inf"),
            str(name or "").casefold())


def sort_schedule(sheet, header: int, last_row: int, columns: dict) -> dict:
    keys = pd.DataFrame(
        [schedule_sort_key(sheet.cell(row, columns["Finish"]).value,
                           sheet.cell(row, columns["Activity Name"]).value)
         for row in range(header + 1, last_row + 1)],
        index=range(header + 1, last_row + 1), columns=["finish", "name"],
    )
    ordered_rows = keys.sort_values(["finish", "name"], kind="stable").index
    mapping = {source: target for target, source in enumerate(ordered_rows, header + 1)}
    moved = {}
    for (row, column), cell in sheet._cells.items():
        cell.row = mapping.get(row, row)
        if cell.hyperlink:
            cell.hyperlink.ref = cell.coordinate
        moved[(cell.row, column)] = cell
    sheet._cells = moved
    dimensions = list(sheet.row_dimensions.items())
    sheet.row_dimensions.clear()
    for row, dimension in dimensions:
        target = mapping.get(row, row)
        dimension.index = target
        sheet.row_dimensions[target] = dimension
    return mapping


def remove_generated_views(book, sheet) -> None:
    for other in list(book):
        if other != sheet and other["A1"].value == STATUS_TITLE and other["B2"].value == sheet.title:
            book.remove(other)
    if CALCULATOR_NAME not in book.defined_names:
        return
    marker = book.defined_names[CALCULATOR_NAME]
    destinations = list(marker.destinations)
    if len(destinations) != 1 or destinations[0][0] != sheet.title:
        raise ValueError("Existing calculator name does not refer to the selected schedule")
    _, reference = destinations[0]
    first_column, first_row, last_column, last_row = openpyxl.utils.range_boundaries(reference)
    if sheet.cell(first_row, first_column).value != "WEEK DELTA CALCULATOR":
        raise ValueError("Existing calculator marker is not recognized")
    sheet._cells = {key: cell for key, cell in sheet._cells.items()
                    if not (first_row <= key[0] <= last_row
                            and first_column <= key[1] <= last_column)}
    for row in range(first_row, last_row + 1):
        if row in sheet.row_dimensions:
            del sheet.row_dimensions[row]
    sheet.data_validations.dataValidation = [
        rule for rule in sheet.data_validations.dataValidation
        if rule.errorTitle != "PTS date input"
    ]
    del book.defined_names[CALCULATOR_NAME]


def create_status_sheet(book, schedule, decisions: pd.Series, columns: dict) -> dict:
    status = book.create_sheet("PTS Status")
    status.sheet_view.showGridLines = False
    status.merge_cells("A1:F1")
    status["A1"] = STATUS_TITLE
    status["A1"].font = Font(name="Calibri", size=18, bold=True, color="FFFFFF")
    status["A1"].fill = PatternFill("solid", fgColor="244C45")
    status.row_dimensions[1].height = 34
    status["A2"], status["B2"] = "Schedule", schedule.title
    missing, found = [], []
    for family, (milestones, disciplines) in RULES.items():
        for milestone in milestones:
            matching_rows = [row for row, item in decisions.items()
                             if item["family"] == family and item["milestone"] == milestone]
            label = f"{family} - {milestone}"
            if not matching_rows:
                missing.append([label, ", ".join(disciplines),
                                "Optional" if milestone.startswith("DTS") else "Required"])
            else:
                for row in matching_rows:
                    found.append([label, ", ".join(disciplines),
                                  schedule.cell(row, columns["Activity ID"]).value,
                                  schedule.cell(row, columns["Start"]).value,
                                  schedule.cell(row, columns["Finish"]).value, row])
    row_number = 4
    for title, headers, entries, colour in (
        (f"NOT FOUND ({len(missing)})", ["Family / Milestone", "Discipline", "Requirement"],
         missing, "9C3030"),
        (f"FOUND ({len(found)})", ["Family / Milestone", "Discipline", "Activity ID", "Start", "Finish", "Schedule Row"],
         found, "244C45"),
    ):
        status.merge_cells(start_row=row_number, start_column=1, end_row=row_number, end_column=6)
        title_cell = status.cell(row_number, 1, title)
        title_cell.font = Font(name="Calibri", size=12, bold=True, color="FFFFFF")
        title_cell.fill = PatternFill("solid", fgColor=colour)
        status.row_dimensions[row_number].height = 26
        row_number += 1
        for column, name in enumerate(headers, 1):
            cell = status.cell(row_number, column, name)
            cell.font = Font(name="Calibri", bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor=colour)
        row_number += 1
        for entry in entries or [["None"]]:
            for column, value in enumerate(entry, 1):
                cell = status.cell(row_number, column, value)
                cell.font = Font(name="Calibri", size=11)
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                family = next((name for name in RULES if str(entry[0]).startswith(name + " - ")), None)
                cell.fill = PatternFill("solid", fgColor=COLOURS[family] if family else "F2F4F3")
                if isinstance(value, datetime):
                    cell.number_format = "DD.MM.YYYY"
                if column == 6 and isinstance(value, int):
                    target = schedule.cell(value, columns["Activity ID"]).coordinate
                    cell.hyperlink = f"#{openpyxl.utils.quote_sheetname(schedule.title)}!{target}"
                    cell.font = Font(name="Calibri", size=11, color="0563C1", underline="single")
            status.row_dimensions[row_number].height = 34
            row_number += 1
        row_number += 2
    for column, width in zip("ABCDEF", (100, 23, 19, 15, 15, 18)):
        status.column_dimensions[column].width = width
    status.freeze_panes = "B6"
    status.print_options.horizontalCentered = True
    status.sheet_properties.pageSetUpPr.fitToPage = True
    status.page_setup.orientation = "landscape"
    status.page_setup.paperSize = status.PAPERSIZE_A3
    status.page_setup.fitToWidth = 1
    status.page_setup.fitToHeight = 0
    status.print_area = f"A1:F{row_number - 1}"
    return {"worksheet": status.title, "missing": [entry[0] for entry in missing],
            "found": len(found)}


def create_week_calculator(book, sheet, last_row: int, columns: dict) -> dict:
    first_row = last_row + 3
    label_column, input_column = columns["Activity Name"], columns["Finish"]
    if label_column >= input_column:
        raise ValueError("Calculator requires Activity Name to appear before Finish")
    border = Border(bottom=Side(style="thin", color="9AA8A3"))
    for offset, label in enumerate(("WEEK DELTA CALCULATOR", "First date", "Second date", "Weeks (second - first)")):
        row_number = first_row + offset
        sheet.row_dimensions[row_number].hidden = False
        sheet.row_dimensions[row_number].height = 27
        for column in (label_column, input_column):
            cell = sheet.cell(row_number, column)
            cell.font = Font(name="Calibri", size=11, bold=offset in (0, 3),
                             color="FFFFFF" if offset == 0 else "243B34")
            cell.fill = PatternFill("solid", fgColor="244C45" if offset == 0 else
                                   "E2F0D9" if offset == 3 else "FFF2CC")
            cell.alignment = Alignment(vertical="center", horizontal="left" if column == label_column else "center")
            cell.border = border
        sheet.cell(row_number, label_column, label)
    upper = sheet.cell(first_row + 1, input_column)
    lower = sheet.cell(first_row + 2, input_column)
    upper.number_format = lower.number_format = "DD.MM.YYYY"
    result = sheet.cell(first_row + 3, input_column)
    result.value = f'=IF(COUNT({upper.coordinate}:{lower.coordinate})=2,({lower.coordinate}-{upper.coordinate})/7,"")'
    result.number_format = '0.00" weeks";[Red]-0.00" weeks";0.00" weeks"'
    validation = DataValidation(type="date", operator="between", formula1="DATE(1900,1,1)",
                                formula2="DATE(9999,12,31)", allow_blank=True)
    validation.errorTitle, validation.error = "PTS date input", "Enter a valid Excel date."
    validation.showErrorMessage = True
    sheet.add_data_validation(validation)
    validation.add(f"{upper.coordinate}:{lower.coordinate}")
    reference = f"${get_column_letter(label_column)}${first_row}:${get_column_letter(input_column)}${first_row + 3}"
    book.defined_names.add(DefinedName(CALCULATOR_NAME,
                                     attr_text=f"{openpyxl.utils.quote_sheetname(sheet.title)}!{reference}"))
    book.calculation.fullCalcOnLoad = True
    book.calculation.forceFullCalc = True
    return {"first_date": upper.coordinate, "second_date": lower.coordinate,
            "weeks": result.coordinate, "formula": result.value}


def process_workbook(path: Path, requested: str | None = None,
                     overwrite: bool = False) -> dict:
    if path.suffix.lower() not in (".xlsx", ".xlsm"):
        raise ValueError("Expected an .xlsx or .xlsm workbook")
    output = path.with_name(path.stem + "_LS" + path.suffix)
    report_path = output.with_suffix(".report.json")
    if not overwrite and (output.exists() or report_path.exists()):
        raise FileExistsError(f"Output already exists: {output}; use --overwrite")
    source_hash = hashlib.sha256(path.read_bytes()).hexdigest()
    book = openpyxl.load_workbook(path, keep_vba=path.suffix.lower() == ".xlsm")
    sheet = select_sheet(book, requested)
    remove_generated_views(book, sheet)
    header, columns, old_classes = detect_headers(sheet)
    other_sheets = {other.title: sheet_fingerprint(other) for other in book if other != sheet}
    original = {(row, column): cell.value for (row, column), cell in sheet._cells.items()
                if cell.value is not None and column not in old_classes.values()}
    width = max(max(column for _, column in original), max(old_classes.values(), default=1))
    last_row = max(row for row, _ in original)
    frame = pd.DataFrame(
        sheet.iter_rows(min_row=header + 1, max_row=last_row, max_col=width, values_only=True),
        index=range(header + 1, last_row + 1), columns=range(1, width + 1),
    )
    decisions = frame[columns["Activity Name"]].fillna("").map(classify)
    report = {"source": str(path), "output": str(output), "worksheet": sheet.title,
              "header_row": header, "counts": {}, "ambiguous": [], "unresolved_dates": [],
              "manual_differences": [], "missing_milestones": [],
              "colours": COLOURS, "overrides": ["Latest visible worksheet", "_LS filename suffix"]}
    for row_number, decision in decisions.items():
        identity = {"row": row_number,
                    "activity_id": str(frame.at[row_number, columns["Activity ID"]] or ""),
                    "activity_name": str(frame.at[row_number, columns["Activity Name"]] or "")}
        if decision["ambiguous"]:
            report["ambiguous"].append(identity)
        for discipline, column in old_classes.items():
            if discipline == "ALL":
                continue
            old_mark = normalize(frame.at[row_number, column]) == "x"
            new_mark = discipline in decision["disciplines"]
            if old_mark != new_mark:
                report["manual_differences"].append(
                    dict(identity, discipline=discipline, manual=old_mark, automatic=new_mark))
    mapping = relocate_columns(sheet, columns["Finish"], old_classes)
    columns = {name: mapping[column] for name, column in columns.items()}
    classification_columns = {name: columns["Finish"] + offset
                              for offset, name in enumerate(DISCIPLINES, start=1)}
    final_width = max(mapping[column] for column in range(1, width + 1) if column in mapping)
    final_width = max(final_width, classification_columns["ALL"])
    for name, column in classification_columns.items():
        cell = sheet.cell(header, column, name)
        cell._style = copy(sheet.cell(header, columns["Activity Name"])._style)
        cell.alignment = Alignment(horizontal="center")
        sheet.column_dimensions[get_column_letter(column)].width = 12
    visible_columns = set(columns.values()) | set(classification_columns.values())
    for column in range(1, final_width + 1):
        sheet.column_dimensions[get_column_letter(column)].hidden = column not in visible_columns
    expected_dates = {}
    for row_number, decision in decisions.items():
        for name in ("Start", "Finish"):
            cell = sheet.cell(row_number, columns[name])
            cleaned, valid = clean_date(cell.value)
            if valid:
                cell.value = cleaned
                if cleaned is not None and cleaned != "":
                    cell.number_format = "DD.MM.YYYY"
            else:
                report["unresolved_dates"].append(
                    {"worksheet": sheet.title, "cell": cell.coordinate, "value": str(cell.value)})
            expected_dates[(row_number, columns[name])] = cleaned
        disciplines = decision["disciplines"]
        for name, column in classification_columns.items():
            marked = bool(disciplines) if name == "ALL" else name in disciplines
            cell = sheet.cell(row_number, column)
            cell.value = "X" if marked else None
            cell.alignment = Alignment(horizontal="center")
        sheet.row_dimensions[row_number].hidden = not bool(disciplines)
        if decision["family"]:
            for row in sheet.iter_rows(min_row=row_number, max_row=row_number, max_col=final_width):
                for cell in row:
                    if disciplines:
                        cell.fill = PatternFill("solid", fgColor=COLOURS[decision["family"]])
                    elif cell.fill.fgColor.type == "rgb" and cell.fill.fgColor.rgb in COLOURS.values():
                        cell.fill = PatternFill()
    for name in DISCIPLINES:
        report["counts"][name] = int(decisions.map(
            lambda item: bool(item["disciplines"]) if name == "ALL"
            else name in item["disciplines"]
        ).sum())
    observed = {(item["family"], item["milestone"]) for item in decisions if item["milestone"]}
    report["missing_milestones"] = [
        f"{family} - {milestone}" for family, (milestones, _) in RULES.items()
        for milestone in milestones if not milestone.startswith("DTS")
        and (family, milestone) not in observed
    ]
    row_mapping = sort_schedule(sheet, header, last_row, columns)
    decisions.index = decisions.index.map(row_mapping)
    decisions = decisions.sort_index()
    report["sort_order"] = ["Finish ascending", "Activity Name ascending (case-insensitive)"]
    report["undated_rows"] = "Blank or unparseable Finish dates sort last"
    for category in ("ambiguous", "manual_differences"):
        for item in report[category]:
            item["source_row"] = item["row"]
            item["row"] = row_mapping[item["row"]]
    for item in report["unresolved_dates"]:
        column_letter, source_row = openpyxl.utils.cell.coordinate_from_string(item["cell"])
        item["source_cell"] = item["cell"]
        item["cell"] = f"{column_letter}{row_mapping[source_row]}"
    sheet.auto_filter = AutoFilter(ref=f"A{header}:{get_column_letter(final_width)}{last_row}")
    sheet.auto_filter.add_filter_column(classification_columns["ALL"] - 1, ["X"])
    sheet.sheet_view.filterMode = True
    sheet.freeze_panes = f"A{header + 1}"
    report["status"] = create_status_sheet(book, sheet, decisions, columns)
    report["calculator"] = create_week_calculator(book, sheet, last_row, columns)
    book.active = book.index(sheet)
    temporary = output.with_name(output.stem + ".validating" + output.suffix)
    try:
        book.save(temporary)
        check = openpyxl.load_workbook(temporary, keep_vba=path.suffix.lower() == ".xlsm")
        try:
            result = check[sheet.title]
            if check.sheetnames != book.sheetnames or check.active.title != sheet.title:
                raise ValueError("Worksheet preservation/activation validation failed")
            for title, fingerprint in other_sheets.items():
                if sheet_fingerprint(check[title]) != fingerprint:
                    raise ValueError(f"Untouched worksheet changed: {title}")
            if detect_headers(result)[2] != classification_columns:
                raise ValueError("Classification layout validation failed")
            for (row, column), value in original.items():
                coordinate = (row, mapping[column])
                expected = expected_dates.get(coordinate, value)
                target = (row_mapping.get(row, row), mapping[column])
                if result.cell(*target).value != expected:
                    raise ValueError(f"Original value changed: {row}, {column}")
            sort_keys = [schedule_sort_key(result.cell(row, columns["Finish"]).value,
                                           result.cell(row, columns["Activity Name"]).value)
                         for row in range(header + 1, last_row + 1)]
            if sort_keys != sorted(sort_keys):
                raise ValueError("Finish date / Activity Name sort validation failed")
            for row_number, decision in decisions.items():
                for name, column in classification_columns.items():
                    marked = bool(decision["disciplines"]) if name == "ALL" else name in decision["disciplines"]
                    if result.cell(row_number, column).value != ("X" if marked else None):
                        raise ValueError(f"Classification validation failed at row {row_number}")
                if bool(result.row_dimensions[row_number].hidden) == bool(decision["disciplines"]):
                    raise ValueError(f"Filter visibility validation failed at row {row_number}")
                for name in ("Start", "Finish"):
                    cell = result.cell(row_number, columns[name])
                    if isinstance(cell.value, datetime) and cell.number_format != "DD.MM.YYYY":
                        raise ValueError(f"Date formatting validation failed: {cell.coordinate}")
                if decision["disciplines"]:
                    for column in range(1, final_width + 1):
                        if result.cell(row_number, column).fill.fgColor.rgb != COLOURS[decision["family"]]:
                            raise ValueError(f"Colour validation failed at row {row_number}")
            if (result.freeze_panes != f"A{header + 1}"
                    or result.auto_filter.ref != sheet.auto_filter.ref
                    or result.auto_filter.filterColumn[0].filters.filter != ["X"]):
                raise ValueError("Filter/freeze validation failed")
            for column in range(1, final_width + 1):
                if result.column_dimensions[get_column_letter(column)].hidden == (column in visible_columns):
                    raise ValueError("Visible column validation failed")
            status = check[report["status"]["worksheet"]]
            expected_missing = report["status"]["missing"]
            if [status.cell(6 + offset, 1).value for offset in range(len(expected_missing))] != expected_missing:
                raise ValueError("Status sheet missing milestones validation failed")
            for coordinate, expected in book[status.title]._cells.items():
                if expected.value is not None:
                    actual = status.cell(*coordinate)
                    if (actual.value != expected.value
                            or actual.number_format != expected.number_format
                            or copy(actual.font) != copy(expected.font)
                            or copy(actual.fill) != copy(expected.fill)
                            or copy(actual.alignment) != copy(expected.alignment)):
                        raise ValueError(
                            f"Status sheet content/style validation failed: {actual.coordinate}; "
                            f"values {actual.value!r} / {expected.value!r}; "
                            f"formats {actual.number_format!r} / {expected.number_format!r}; "
                            f"font={copy(actual.font) == copy(expected.font)}, "
                            f"fill={copy(actual.fill) == copy(expected.fill)}, "
                            f"alignment={copy(actual.alignment) == copy(expected.alignment)}")
            calculator = report["calculator"]
            if result[calculator["weeks"]].value != calculator["formula"]:
                raise ValueError("Week calculator formula validation failed")
            if (result[calculator["first_date"]].value is not None
                    or result[calculator["second_date"]].value is not None
                    or result[calculator["first_date"]].row <= last_row
                    or result.row_dimensions[result[calculator["weeks"]].row].hidden
                    or not check.calculation.fullCalcOnLoad):
                raise ValueError("Week calculator placement/input validation failed")
        finally:
            check.close()
        if hashlib.sha256(path.read_bytes()).hexdigest() != source_hash:
            raise ValueError("Source workbook changed during processing")
        temporary.replace(output)
        report["validated"] = True
        report_path.write_text(json.dumps(report, indent=2), encoding="utf-8")
    finally:
        book.close()
        if temporary.exists():
            temporary.unlink()
    return report


def inspect_workbook(path: Path) -> None:
    book = openpyxl.load_workbook(path, read_only=True, data_only=False)
    try:
        print(f"\nFILE: {path.name}")
        for sheet in book:
            print(f"SHEET: {sheet.title!r}, {sheet.sheet_state}, "
                  f"{sheet.max_row} rows x {sheet.max_column} columns")
            if sheet.sheet_state != "visible" or "Apr" in sheet.title:
                continue
            frame = pd.DataFrame(sheet.iter_rows(max_col=25, values_only=True))
            print("HEADERS:", frame.iloc[0].dropna().to_dict())
            marks = frame.iloc[:, 19:21].apply(
                lambda column: column.astype(str).str.upper().eq("X")
            ).any(axis=1)
            print("MARKS:")
            print(frame.loc[marks, [0, 1, 19, 20]].to_string(index=True))
            colours = Counter()
            for row in sheet.iter_rows(min_row=2, max_col=25):
                fills = tuple(
                    (cell.column_letter, cell.fill.fgColor.type,
                     str(cell.fill.fgColor.index))
                    for cell in row if cell.fill is not None
                    and cell.fill.patternType == "solid"
                )
                if fills:
                    colours[fills] += 1
                    print("COLOUR:", row[0].row, row[0].value, row[1].value, fills)
            print("COLOUR COUNTS:", colours)
    finally:
        book.close()


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("files", nargs="*", type=Path)
    parser.add_argument("--inspect", action="store_true")
    parser.add_argument("--sheet", help="Explicit visible worksheet if date detection is ambiguous")
    parser.add_argument("--overwrite", action="store_true", help="Replace existing LS outputs only")
    args = parser.parse_args()
    files = args.files or sorted(path for path in Path(__file__).parent.glob("*.xls*")
                                if path.suffix.lower() in (".xlsx", ".xlsm")
                                and not path.stem.endswith(("_LS", ".validating"))
                                and not path.name.startswith("~$"))
    for path in files:
        if args.inspect:
            inspect_workbook(path)
        else:
            report = process_workbook(path, args.sheet, args.overwrite)
            print(json.dumps({key: report[key] for key in (
                "output", "worksheet", "header_row", "counts", "missing_milestones")}, indent=2))
            print(f"Ambiguous: {len(report['ambiguous'])}; unresolved dates: "
                  f"{len(report['unresolved_dates'])}; manual differences: "
                  f"{len(report['manual_differences'])}; reopened and validated.")


if __name__ == "__main__":
    main()