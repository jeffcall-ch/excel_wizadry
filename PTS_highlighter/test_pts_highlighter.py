"""Focused classification and workbook preservation checks."""

from datetime import datetime
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

import openpyxl

from pts_highlighter import (
    COLOURS, ENGINEERING, RULES, classify, clean_date, process_workbook, select_sheet,
    sort_schedule,
)


class ScheduleTests(unittest.TestCase):
    def test_every_approved_rule(self):
        for family, (milestones, disciplines) in RULES.items():
            for milestone in milestones:
                with self.subTest(family=family, milestone=milestone):
                    self.assertEqual(classify(f"{family} - {milestone}")["disciplines"], disciplines)

    def test_guardrails(self):
        for name in (
            "Lot General Piping Pipe Material",
            "Lot General Piping Pipe Material - Engineering & Manufacturing (by Supplier)",
            "Lot General Piping Pipe Material - TSD Documents complete / RFQ",
            "Lot General Piping Engineering - Detail Engineering - ISO and BOM for Erection - Issued",
            "Lot 5 - Erection General Piping - Contractor On Site Mobilisation",
            "Lot 6 - Erection General Piping - BOQ 1 (Doc's for RfQ) - Issued from Engineering",
            "Other Lot - Purchase Order",
            "Lot General Piping Valves - Purchase Order Cancelled",
        ):
            with self.subTest(name=name):
                self.assertFalse(classify(name)["disciplines"])

    def test_normalization_and_small_typo(self):
        result = classify("LOT General Piping Pipe Material: Tender Due Daate")
        self.assertEqual(result["disciplines"], ("Piping",))
        result = classify("Lot General Piping Pipe Material - Technical & Commercial Comparison Pre TCC")
        self.assertEqual(result["milestone"], "Technical & Commercial Comparison Pre-TCC")
        self.assertEqual(classify("Lot General Piping Hangers - DTS (S)")["milestone"], "DTS (S)")

    def test_dates(self):
        for text in ("22.09.2026 A", "*22.09.2026", "22/09/2026*", " 22-09-2026 "):
            self.assertEqual(clean_date(text), (datetime(2026, 9, 22), True))
        for text in ("31.02.2026 A", "not a date", "01.02.2026 / 02.02.2026"):
            self.assertEqual(clean_date(text), (text, False))
        self.assertEqual(clean_date(None), (None, True))

    def test_latest_visible(self):
        book = openpyxl.Workbook()
        book.active.title = "DD25Jul2026 Update"
        book.create_sheet("DD25Sep2026 Update")
        book.create_sheet("DD25Dec2026 Update").sheet_state = "hidden"
        self.assertEqual(select_sheet(book).title, "DD25Sep2026 Update")
        book.create_sheet("Unknown")
        with self.assertRaises(ValueError):
            select_sheet(book)

    def test_processing_and_repeat_run(self):
        with TemporaryDirectory() as directory:
            source = Path(directory) / "schedule.xlsx"
            book = openpyxl.Workbook()
            sheet = book.active
            sheet.title = "DD25Sep2026 Update"
            sheet.append(["Schedule"])
            sheet.append(["Activity ID", "Activity Name", "BL Project Start", "Start", "Finish", "Piping ", "Hanger", "All"])
            sheet.append(["ENG1", "Lot General Piping Engineering - " + ENGINEERING[0],
                          "baseline unchanged", "*22.09.2026 A", None, "x", None, '=IF(F3="x","X","")'])
            sheet.append(["BAD1", "Lot General Piping Pipe Material - Engineering & Manufacturing (by Supplier)",
                          None, "bad date", "23/09/2026*", "x"])
            sheet.append(["VAL1", "Lot General Piping Valves - Purchase Order",
                          None, "01/10/2026", "02/10/2026"])
            sheet.cell(3, 16373).fill = openpyxl.styles.PatternFill("solid", fgColor="FFEEEEEE")
            archive = book.create_sheet("Archive")
            archive.sheet_state = "hidden"
            archive["A1"] = "Preserved"
            book.save(source)
            report = process_workbook(source)
            self.assertTrue(report["validated"])
            self.assertEqual(report["counts"], {"Piping": 1, "Hangers": 1, "Valves": 1, "ALL": 2})
            self.assertEqual(len(report["unresolved_dates"]), 1)
            output = Path(report["output"])
            check = openpyxl.load_workbook(output)
            result = check.active
            self.assertEqual([result.cell(2, column).value for column in range(6, 10)],
                             ["Piping", "Hangers", "Valves", "ALL"])
            self.assertEqual([result.cell(row, 1).value for row in range(3, 6)],
                             ["BAD1", "VAL1", "ENG1"])
            self.assertEqual(result["C5"].value, "baseline unchanged")
            self.assertEqual(result["D5"].value, datetime(2026, 9, 22))
            self.assertEqual(result["B5"].fill.fgColor.rgb, COLOURS["Lot General Piping Engineering"])
            self.assertEqual(check["Archive"].sheet_state, "hidden")
            self.assertEqual(result.cell(5, 16374).fill.fgColor.rgb, "FFEEEEEE")
            status = check[report["status"]["worksheet"]]
            self.assertEqual(status["A4"].value, "NOT FOUND (29)")
            missing = report["status"]["missing"]
            self.assertIn("Lot General Piping Pipe Material - DTS (S)", missing)
            self.assertEqual([status.cell(row, 1).value for row in range(6, 6 + len(missing))], missing)
            self.assertEqual(report["status"]["found"], 2)
            calculator = report["calculator"]
            self.assertEqual(calculator["first_date"], "E9")
            self.assertEqual(calculator["second_date"], "E10")
            self.assertEqual(result["E11"].value, '=IF(COUNT(E9:E10)=2,(E10-E9)/7,"")')
            self.assertEqual(result["E9"].number_format, "DD.MM.YYYY")
            self.assertGreater(result["E9"].row, 5)
            self.assertFalse(result.row_dimensions[11].hidden)
            self.assertTrue(check.calculation.fullCalcOnLoad)
            self.assertEqual(str(result.data_validations.dataValidation[0].sqref), "E9:E10")
            check.close()
            with self.assertRaises(FileExistsError):
                process_workbook(source)
            second = process_workbook(output)
            self.assertEqual(second["counts"], report["counts"])
            self.assertEqual(second["status"], report["status"])
            self.assertEqual(second["calculator"], report["calculator"])

    def test_finish_sort_with_name_ties_and_undated_rows(self):
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.append(["Activity ID", "Activity Name", "Finish"])
        sheet.append(["LATE", "Alpha", datetime(2026, 10, 2)])
        sheet.append(["ZULU", "Zulu", datetime(2026, 10, 1)])
        sheet.append(["EMPTY", "Undated B", None])
        sheet.append(["ALPHA1", "alpha", datetime(2026, 10, 1)])
        sheet.append(["INVALID", "Undated A", "bad date"])
        sheet.append(["ALPHA2", "alpha", datetime(2026, 10, 1)])
        sheet.row_dimensions[5].height = 42
        sheet.row_dimensions[5].hidden = True
        sheet["A5"].hyperlink = "https://example.com"
        mapping = sort_schedule(sheet, 1, 7, {"Activity Name": 2, "Finish": 3})
        self.assertEqual([sheet.cell(row, 1).value for row in range(2, 8)],
                         ["ALPHA1", "ALPHA2", "ZULU", "LATE", "INVALID", "EMPTY"])
        self.assertEqual(mapping[5], 2)
        self.assertEqual(sheet.row_dimensions[2].height, 42)
        self.assertTrue(sheet.row_dimensions[2].hidden)
        self.assertEqual(sheet["A2"].hyperlink.ref, "A2")
        self.assertEqual(sheet["A1"].value, "Activity ID")


if __name__ == "__main__":
    unittest.main()