"""Walk a folder of pipe isometric PDF drawings and export a formatted
Excel table listing KKS, revision and system parsed from the file names.

Expected file name pattern (example):
    TB020-KVI-1QFB99BR240_0.0_Pipe-Isometric-Drawing-Service_Air_BOP.pdf

    TB020-KVI-                -> project/discipline prefix, dropped
    1QFB99BR240                -> KKS (kept)
    0.0                         -> revision, kept as number with 1 decimal
    Pipe-Isometric-Drawing      -> document type, dropped
    Service_Air_BOP             -> system name, kept
"""
from __future__ import annotations

import argparse
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

# Matches: <prefix>-<prefix>-<KKS>_<revision>_Pipe-Isometric-Drawing-<System>.pdf
FILENAME_PATTERN = re.compile(
    r"^.+?-.+?-(?P<kks>[^_]+)_(?P<revision>\d+(?:\.\d+)?)_"
    r"Pipe-Isometric-Drawing-(?P<system>.+)$",
    re.IGNORECASE,
)


def parse_filename(stem: str) -> tuple[str, float, str] | None:
    """Return (kks, revision, system) parsed from a PDF file stem, or None if it doesn't match."""
    match = FILENAME_PATTERN.match(stem)
    if not match:
        return None
    kks = match.group("kks")
    revision = float(match.group("revision"))
    system = match.group("system").rstrip("_")
    return kks, revision, system


def collect_rows(folder: Path) -> tuple[list[tuple[str, str, float, str]], list[str]]:
    """Return (rows, unmatched_filenames) for all PDFs found under folder (recursive)."""
    rows: list[tuple[str, str, float, str]] = []
    unmatched: list[str] = []
    for pdf_path in sorted(folder.rglob("*.pdf")):
        parsed = parse_filename(pdf_path.stem)
        if parsed is None:
            unmatched.append(pdf_path.name)
            continue
        kks, revision, system = parsed
        rows.append((pdf_path.name, kks, revision, system))
    return rows, unmatched


def write_workbook(rows: list[tuple[str, str, float, str]], output_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Isometrics"

    headers = ["File Name", "KKS", "Revision", "System"]
    ws.append(headers)
    for row in rows:
        ws.append(row)

    # Revision column formatted as a number with one decimal digit.
    revision_col = headers.index("Revision") + 1
    for row_idx in range(2, ws.max_row + 1):
        ws.cell(row=row_idx, column=revision_col).number_format = "0.0"

    # Auto-size columns.
    for col_idx, header in enumerate(headers, start=1):
        max_len = max([len(header)] + [len(str(r[col_idx - 1])) for r in rows]) if rows else len(header)
        ws.column_dimensions[get_column_letter(col_idx)].width = max_len + 2

    last_row = ws.max_row
    last_col_letter = get_column_letter(len(headers))
    table_ref = f"A1:{last_col_letter}{last_row}"
    table = Table(displayName="IsometricsTable", ref=table_ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium9",
        showRowStripes=True,
    )
    ws.add_table(table)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "folder",
        nargs="?",
        default=r"C:\temp\0_Thameside_isometrics_All\Tilbury\Att.1b",
        help="Folder containing the PDF isometric drawings.",
    )
    parser.add_argument(
        "--output-name",
        default="isometric_filenames.xlsx",
        help="Name of the generated xlsx file.",
    )
    args = parser.parse_args()

    folder = Path(args.folder)
    if not folder.is_dir():
        raise SystemExit(f"Folder not found: {folder}")

    rows, unmatched = collect_rows(folder)
    output_path = folder / "xlsx" / args.output_name
    write_workbook(rows, output_path)

    print(f"Wrote {len(rows)} rows to {output_path}")
    if unmatched:
        print(f"{len(unmatched)} file(s) did not match the expected pattern:")
        for name in unmatched:
            print(f"  - {name}")


if __name__ == "__main__":
    main()
