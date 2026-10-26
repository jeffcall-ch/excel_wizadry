from __future__ import annotations

import argparse
import configparser
from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterable

import openpyxl
import pandas as pd
from openpyxl.worksheet.worksheet import Worksheet

DEFAULT_CONFIG_FILENAME = "system_batching.ini"

REQUIRED_CSV_COLUMNS = [
    "support_structure_su",
    "resolved_system_first_in_priority",
]
REQUIRED_WEIGHT_COLUMNS = ["Support Structure", "Total Weight [kg]"]

DETAIL_COLUMN_WIDTHS = {
    "A": 22.00,
    "B": 21.00,
    "C": 16.00,
    "D": 12.00,
    "E": 23.00,
    "F": 30.00,
    "G": 18.00,
    "H": 15.00,
}

SUMMARY_COLUMN_WIDTHS = {
    "A": 13.00,
    "B": 16.00,
    "C": 36.00,
    "D": 13.00,
    "E": 10.00,
    "F": 17.00,
    "G": 16.00,
    "H": 16.00,
    "I": 18.00,
    "J": 17.00,
}


@dataclass
class Settings:
    resolved_csv: Path
    original_file: Path
    original_sheet: str
    output_file: Path
    batch_count: int
    target_weight_kg: float
    tolerance: float
    system_priority: list[str]
    support_system_overrides: dict[str, str]


@dataclass
class BatchBin:
    batch_number: int
    weight_kg: float = 0.0
    support_codes: list[str] = field(default_factory=list)
    systems: set[str] = field(default_factory=set)


def normalize_kks(value: str) -> str:
    token = str(value).strip().upper()
    if token.startswith("/"):
        token = token[1:]
    return token


def parse_priority_list(value: str) -> list[str]:
    items = [part.strip() for part in str(value).split(",") if part.strip()]
    deduped: list[str] = []
    seen: set[str] = set()
    for item in items:
        key = item.lower()
        if key in seen:
            continue
        seen.add(key)
        deduped.append(item)
    return deduped


def resolve_path(base_dir: Path, raw_path: str) -> Path:
    candidate = Path(raw_path).expanduser()
    if candidate.is_absolute():
        return candidate
    return (base_dir / candidate).resolve()


def read_settings(config_path: Path) -> Settings:
    parser = configparser.ConfigParser()
    if not parser.read(config_path, encoding="utf-8"):
        raise FileNotFoundError(f"Config file not found: {config_path}")

    base_dir = config_path.parent

    resolved_csv = resolve_path(
        base_dir,
        parser.get("paths", "resolved_csv", fallback="su_system_sequence_resolved.csv"),
    )
    original_file = resolve_path(
        base_dir,
        parser.get(
            "paths",
            "original_file",
            fallback="CA100-KVI-50296373_2.0 - BoM General Hangers and support material ORIGINAL DATA .xlsm",
        ),
    )
    output_file = resolve_path(
        base_dir,
        parser.get("paths", "output_file", fallback="su_system_batched.xlsx"),
    )

    original_sheet = parser.get(
        "paths",
        "original_sheet",
        fallback="Per_Support_Structure",
    )

    batch_count = parser.getint("batching", "batch_count", fallback=5)
    target_weight_kg = parser.getfloat("batching", "target_weight_kg", fallback=4000.0)
    tolerance = parser.getfloat("batching", "tolerance", fallback=0.25)

    priority = parse_priority_list(
        parser.get(
            "systems",
            "priority",
            fallback=(
                "Instrument air,Service air,Compressed air,Nitrogen,"
                "Process Water,Towns Water,Bottom ash,Drain water,"
                "Blow down,Exhaust Air,Condensate drains,HP water,"
                "Fuel oil,Ammonia water"
            ),
        )
    )

    if batch_count <= 0:
        raise ValueError("batch_count must be a positive integer")
    if target_weight_kg <= 0:
        raise ValueError("target_weight_kg must be > 0")
    if not (0.0 < tolerance < 1.0):
        raise ValueError("tolerance must be between 0 and 1")
    if not priority:
        raise ValueError("systems.priority in config cannot be empty")

    priority_lookup = {system.lower(): system for system in priority}
    overrides: dict[str, str] = {}
    if parser.has_section("overrides"):
        for raw_support, raw_system in parser.items("overrides"):
            support_code = normalize_kks(raw_support)
            if not support_code.endswith("/SU"):
                support_code = f"{support_code}/SU"

            system_text = str(raw_system).strip()
            canonical_system = priority_lookup.get(system_text.lower(), system_text)
            if canonical_system:
                overrides[support_code] = canonical_system

    return Settings(
        resolved_csv=resolved_csv,
        original_file=original_file,
        original_sheet=original_sheet,
        output_file=output_file,
        batch_count=batch_count,
        target_weight_kg=target_weight_kg,
        tolerance=tolerance,
        system_priority=priority,
        support_system_overrides=overrides,
    )


def parse_args() -> argparse.Namespace:
    script_dir = Path(__file__).resolve().parent
    parser = argparse.ArgumentParser(
        description=(
            "Create 5 weight-limited shipping batches by system priority, "
            "without splitting systems unless required by max batch weight."
        )
    )
    parser.add_argument(
        "--config",
        type=Path,
        default=script_dir / DEFAULT_CONFIG_FILENAME,
        help="Path to INI configuration file.",
    )
    return parser.parse_args()


def find_header_row(
    workbook_path: Path,
    sheet_name: str,
    required_columns: Iterable[str],
    max_scan_rows: int = 50,
) -> int:
    workbook = openpyxl.load_workbook(
        workbook_path,
        data_only=True,
        read_only=True,
        keep_vba=True,
    )
    worksheet = workbook[sheet_name]
    required = {str(col).strip() for col in required_columns}

    for row_idx, row in enumerate(
        worksheet.iter_rows(
            min_row=1,
            max_row=max_scan_rows,
            values_only=True,
        ),
        start=1,
    ):
        values = {str(value).strip() for value in row if value not in (None, "")}
        if required.issubset(values):
            return row_idx

    raise ValueError(
        f"Could not detect header row with required columns {sorted(required)} "
        f"in sheet '{sheet_name}' of file {workbook_path}"
    )


def load_weight_lookup(original_file: Path, sheet_name: str) -> pd.DataFrame:
    header_row = find_header_row(
        workbook_path=original_file,
        sheet_name=sheet_name,
        required_columns=REQUIRED_WEIGHT_COLUMNS,
    )

    data = pd.read_excel(
        original_file,
        sheet_name=sheet_name,
        engine="openpyxl",
        header=header_row - 1,
    )

    missing = [col for col in REQUIRED_WEIGHT_COLUMNS if col not in data.columns]
    if missing:
        raise ValueError(f"Missing required columns in weight sheet: {missing}")

    weights = data.loc[:, REQUIRED_WEIGHT_COLUMNS].copy()
    weights["Support Structure"] = (
        weights["Support Structure"].astype(str).str.strip().map(normalize_kks)
    )
    weights = weights[weights["Support Structure"].str.endswith("/SU")].copy()
    weights["Total Weight [kg]"] = pd.to_numeric(
        weights["Total Weight [kg]"],
        errors="coerce",
    ).fillna(0.0)

    return (
        weights.groupby("Support Structure", as_index=False)["Total Weight [kg]"]
        .sum()
        .rename(
            columns={
                "Support Structure": "support_structure_su",
                "Total Weight [kg]": "weight_kg",
            }
        )
    )


def load_scope(
    resolved_csv: Path,
    weight_lookup: pd.DataFrame,
    system_priority: list[str],
    support_system_overrides: dict[str, str],
) -> pd.DataFrame:
    scope = pd.read_csv(resolved_csv)

    missing = [col for col in REQUIRED_CSV_COLUMNS if col not in scope.columns]
    if missing:
        raise ValueError(f"Missing required columns in resolved CSV: {missing}")

    scope = scope.copy()
    scope["support_structure_su"] = (
        scope["support_structure_su"].astype(str).str.strip().map(normalize_kks)
    )
    scope = scope.drop_duplicates(subset=["support_structure_su"], keep="first").copy()

    scope = scope.merge(weight_lookup, on="support_structure_su", how="left")
    scope["weight_kg"] = scope["weight_kg"].fillna(0.0)

    scope["resolved_system_first_in_priority"] = (
        scope["resolved_system_first_in_priority"].fillna("").astype(str).str.strip()
    )
    scope.loc[
        scope["resolved_system_first_in_priority"].str.lower() == "nan",
        "resolved_system_first_in_priority",
    ] = ""
    scope.loc[
        scope["resolved_system_first_in_priority"] == "",
        "resolved_system_first_in_priority",
    ] = "UNKNOWN"

    if support_system_overrides:
        override_series = scope["support_structure_su"].map(support_system_overrides)
        scope["resolved_system_first_in_priority"] = override_series.fillna(
            scope["resolved_system_first_in_priority"]
        )

    priority_map = {system: idx for idx, system in enumerate(system_priority, start=1)}
    scope["system_rank"] = scope["resolved_system_first_in_priority"].map(priority_map)
    scope["system_rank"] = scope["system_rank"].fillna(999).astype(int)

    for optional_col in [
        "original_owner_system",
        "all_systems_mapped_priority_order",
        "system_changed_for_schedule_risk",
    ]:
        if optional_col not in scope.columns:
            scope[optional_col] = ""
        else:
            scope[optional_col] = scope[optional_col].fillna("").astype(str).str.strip()
            scope.loc[scope[optional_col].str.lower() == "nan", optional_col] = ""

    return scope


def assign_supports_to_batches_sequential(
    scope: pd.DataFrame,
    batch_count: int,
    max_batch_weight: float,
    system_priority: list[str],
) -> tuple[list[BatchBin], dict[str, int]]:
    batches = [BatchBin(batch_number=i + 1) for i in range(batch_count)]
    support_to_batch: dict[str, int] = {}
    current_batch_idx = 0

    priority_map = {system: idx for idx, system in enumerate(system_priority, start=1)}

    ordered = scope.copy()
    ordered["_system_sort"] = ordered["resolved_system_first_in_priority"].map(priority_map).fillna(999)
    ordered = ordered.sort_values(
        ["_system_sort", "resolved_system_first_in_priority", "support_structure_su"],
        ascending=[True, True, True],
    )

    grouped = ordered.groupby("resolved_system_first_in_priority", sort=False)

    for system, system_rows in grouped:
        rows = system_rows.sort_values("support_structure_su").reset_index(drop=True)
        system_total = float(rows["weight_kg"].sum())

        if rows["weight_kg"].max() > max_batch_weight:
            heavy = rows[rows["weight_kg"] > max_batch_weight]["support_structure_su"].tolist()
            raise ValueError(
                "At least one /SU exceeds max batch weight, cannot batch without violating limit: "
                f"{heavy}"
            )

        # Keep system whole whenever it fits in one batch.
        if system_total <= max_batch_weight:
            current_batch = batches[current_batch_idx]
            if (
                current_batch.weight_kg + system_total > max_batch_weight
                and current_batch_idx < (batch_count - 1)
            ):
                current_batch_idx += 1
                current_batch = batches[current_batch_idx]

            if current_batch.weight_kg + system_total <= max_batch_weight:
                supports = rows["support_structure_su"].tolist()
                current_batch.weight_kg += system_total
                current_batch.support_codes.extend(supports)
                current_batch.systems.add(system)
                for support_code in supports:
                    support_to_batch[support_code] = current_batch.batch_number
                continue

        # Split across batches only when unavoidable by max batch limit.
        for _, row in rows.iterrows():
            support_code = str(row["support_structure_su"])
            weight = float(row["weight_kg"])

            while batches[current_batch_idx].weight_kg + weight > max_batch_weight:
                if current_batch_idx >= (batch_count - 1):
                    raise ValueError(
                        "Cannot place all supports within batch max weight and batch_count. "
                        "Increase batch_count or tolerance."
                    )
                current_batch_idx += 1

            batch = batches[current_batch_idx]
            batch.weight_kg += weight
            batch.support_codes.append(support_code)
            batch.systems.add(system)
            support_to_batch[support_code] = batch.batch_number

    return batches, support_to_batch


def build_output_tables(
    scope: pd.DataFrame,
    support_to_batch: dict[str, int],
    system_priority: list[str],
    target_weight_kg: float,
    tolerance: float,
) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    detail = scope.copy()
    detail["Batch number"] = detail["support_structure_su"].map(support_to_batch)
    if detail["Batch number"].isna().any():
        unresolved = detail[detail["Batch number"].isna()]["support_structure_su"].tolist()
        raise ValueError(f"Unassigned /SU supports: {unresolved}")

    priority_map = {system: idx for idx, system in enumerate(system_priority, start=1)}
    detail["_system_sort"] = detail["resolved_system_first_in_priority"].map(priority_map).fillna(999)

    detail = detail.sort_values(
        ["Batch number", "_system_sort", "resolved_system_first_in_priority", "support_structure_su"],
        ascending=[True, True, True, True],
    ).reset_index(drop=True)

    detail_out = detail.rename(
        columns={
            "support_structure_su": "KKS code",
            "resolved_system_first_in_priority": "System",
            "weight_kg": "KKS Total Weight [kg]",
            "original_owner_system": "Original owner system",
            "all_systems_mapped_priority_order": "All systems",
            "system_changed_for_schedule_risk": "System changed for schedule risk",
        }
    )[
        [
            "KKS code",
            "System",
            "KKS Total Weight [kg]",
            "Batch number",
            "Original owner system",
            "All systems",
            "System changed for schedule risk",
        ]
    ]

    min_weight = 0.0
    max_weight = target_weight_kg * (1.0 + tolerance)

    def systems_text(values: pd.Series) -> str:
        systems = sorted(set(values), key=lambda s: priority_map.get(s, 999))
        return ", ".join(systems)

    summary = (
        detail_out.groupby("Batch number", as_index=False)
        .agg(
            Total_Weight_kg=("KKS Total Weight [kg]", "sum"),
            Systems=("System", systems_text),
            System_Count=("System", "nunique"),
            KKS_Count=("KKS code", "nunique"),
        )
        .sort_values("Batch number")
        .reset_index(drop=True)
    )

    summary["Target_Weight_kg"] = target_weight_kg
    summary["Min_Allowed_kg"] = min_weight
    summary["Max_Allowed_kg"] = max_weight
    summary["Delta_vs_Target_kg"] = summary["Total_Weight_kg"] - target_weight_kg
    summary["Within_Tolerance"] = summary["Total_Weight_kg"] <= max_weight

    system_summary = (
        detail_out.groupby(["Batch number", "System"], as_index=False)["KKS Total Weight [kg]"]
        .sum()
        .assign(_system_sort=lambda d: d["System"].map(priority_map).fillna(999))
        .sort_values(["Batch number", "_system_sort", "System"])
        .drop(columns=["_system_sort"])
        .rename(columns={"KKS Total Weight [kg]": "System_Total_Weight_kg"})
    )

    return detail_out, summary, system_summary


def apply_detail_formatting(worksheet: Worksheet, detail_rows: int) -> None:
    last_data_row = detail_rows + 1
    subtotal_formula = f"=SUBTOTAL(109,C2:C{last_data_row})"

    worksheet.auto_filter.ref = f"A1:G{last_data_row}"
    worksheet.freeze_panes = "A2"

    for column, width in DETAIL_COLUMN_WIDTHS.items():
        worksheet.column_dimensions[column].width = width

    worksheet["I1"] = "Total weight"
    worksheet["I1"].style = worksheet["A1"].style

    worksheet["J1"] = subtotal_formula
    worksheet["J1"].number_format = "0.00"

    bottom_total_row = last_data_row + 3
    worksheet.cell(row=bottom_total_row, column=3, value=subtotal_formula)
    worksheet.cell(row=bottom_total_row, column=3).style = worksheet["C2"].style


def apply_summary_formatting(worksheet: Worksheet) -> None:
    for column, width in SUMMARY_COLUMN_WIDTHS.items():
        worksheet.column_dimensions[column].width = width


def write_output(
    output_file: Path,
    detail: pd.DataFrame,
    summary: pd.DataFrame,
    system_summary: pd.DataFrame,
) -> None:
    with pd.ExcelWriter(output_file, engine="openpyxl") as writer:
        detail.to_excel(writer, sheet_name="Detail", index=False)
        summary.to_excel(writer, sheet_name="Summary", index=False, startrow=0)
        system_summary.to_excel(
            writer,
            sheet_name="Summary",
            index=False,
            startrow=len(summary) + 3,
        )

        apply_detail_formatting(worksheet=writer.sheets["Detail"], detail_rows=len(detail))
        apply_summary_formatting(worksheet=writer.sheets["Summary"])


def main() -> None:
    args = parse_args()
    config_path = args.config.resolve()

    settings = read_settings(config_path=config_path)

    if not settings.resolved_csv.exists():
        raise FileNotFoundError(f"Resolved CSV not found: {settings.resolved_csv}")
    if not settings.original_file.exists():
        raise FileNotFoundError(f"Original workbook not found: {settings.original_file}")

    weight_lookup = load_weight_lookup(
        original_file=settings.original_file,
        sheet_name=settings.original_sheet,
    )
    scope = load_scope(
        resolved_csv=settings.resolved_csv,
        weight_lookup=weight_lookup,
        system_priority=settings.system_priority,
        support_system_overrides=settings.support_system_overrides,
    )

    max_batch_weight = settings.target_weight_kg * (1.0 + settings.tolerance)

    batches, support_to_batch = assign_supports_to_batches_sequential(
        scope=scope,
        batch_count=settings.batch_count,
        max_batch_weight=max_batch_weight,
        system_priority=settings.system_priority,
    )

    detail, summary, system_summary = build_output_tables(
        scope=scope,
        support_to_batch=support_to_batch,
        system_priority=settings.system_priority,
        target_weight_kg=settings.target_weight_kg,
        tolerance=settings.tolerance,
    )

    write_output(
        output_file=settings.output_file,
        detail=detail,
        summary=summary,
        system_summary=system_summary,
    )

    split_systems = (
        detail.groupby("System", as_index=False)["Batch number"].nunique().query("`Batch number` > 1")
    )

    print(f"Config file: {config_path}")
    print(f"Resolved CSV: {settings.resolved_csv}")
    print(f"Original workbook: {settings.original_file}")
    print(f"Original sheet: {settings.original_sheet}")
    print(f"Output file: {settings.output_file}")
    print(f"Rows batched (/SU): {len(detail)}")
    print(f"Batch count: {settings.batch_count}")
    print(
        "Target/min/max [kg]: "
        f"{settings.target_weight_kg:.2f} / "
        f"{0.0:.2f} / "
        f"{max_batch_weight:.2f}"
    )
    print("Batch totals [kg]:")
    for _, row in summary.iterrows():
        print(
            f"  Batch {int(row['Batch number'])}: "
            f"{row['Total_Weight_kg']:.2f} "
            f"(within tolerance: {bool(row['Within_Tolerance'])})"
        )
    print(f"Systems split across multiple batches: {len(split_systems)}")


if __name__ == "__main__":
    main()
